"""Build agent-facing mob spreadsheets + diagrams into docs/zero-g-tweaks-bundle/sheets/mobs/.

Sources: mobs_data.py / mobs_food.py (design text, same as the PNG reference sheets),
mob_specs.py (implementation specs), and the repo itself (loot tables, GeckoLib models and
animations, textures, spawn-egg colours, pending spawn biome modifiers, recipes).
Re-run after changing any of those:  python3 mob_docs.py
"""
import csv, glob, json, os, re, subprocess, sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import mobs_data, mobs_food
from mob_specs import S

REPO = '/home/claude/zerog-tweaks'
RES = f'{REPO}/src/main/resources'
A = f'{RES}/assets/zerog_tweaks'
OUT = f'{REPO}/docs/zero-g-tweaks-bundle/sheets/mobs'
DATA, DIA = f'{OUT}/data', f'{OUT}/diagrams'
os.makedirs(DATA, exist_ok=True); os.makedirs(f'{DIA}/states', exist_ok=True)

DESIGN = {**mobs_data.MOBS, **mobs_food.M2}
ORDER = sorted(os.listdir(OUT))
SHEET = {re.sub(r'^\d+_|\.png$', '', f): f for f in ORDER if f.endswith('.png')}
assert set(DESIGN) == set(S) == set(SHEET), (set(DESIGN) ^ set(S), set(SHEET) ^ set(S))
MOBS = sorted(S, key=lambda k: SHEET[k])

LANG = json.load(open(f'{A}/lang/en_us.json'))
def item_name(i):
    ns, n = i.split(':')
    return LANG.get(f'item.{ns}.{n}') or LANG.get(f'block.{ns}.{n}') or n.replace('_', ' ').title()

eggs = {}
ej = open(f'{REPO}/docs/zero-g-tweaks-bundle/java/item/ZGSpawnEggs.java').read()
for k, a, b in re.findall(r'egg\("(\w+)", ZGEntities\.\w+, 0x([0-9A-Fa-f]{6}), 0x([0-9A-Fa-f]{6})\)', ej):
    eggs[k] = (f'#{a.upper()}', f'#{b.upper()}')

def first_int(s):
    m = re.search(r'\d+', s or ''); return int(m.group()) if m else 0

def hitbox(info):
    m = re.search(r'hitbox ([\d.]+) × ([\d.]+)(?: × ([\d.]+))?', info)
    w, h, l = float(m.group(1)), float(m.group(2)), float(m.group(3)) if m.group(3) else None
    return w, h, l

# ------------------------------------------------------------------ loot -> drops
def count_of(fns):
    lo = hi = 1
    for f in fns or []:
        if f['function'] == 'minecraft:set_count':
            c = f['count']
            if isinstance(c, (int, float)): lo = hi = c
            else: lo, hi = c.get('min', 1), c.get('max', 1)
    return lo, hi

def drops(mob):
    j = json.load(open(f'{RES}/data/zerog_tweaks/loot_table/entities/{mob}.json'))
    out = []
    for p in j['pools']:
        pc = p.get('conditions', [])
        for e in p['entries']:
            if e.get('type') != 'minecraft:item': continue
            fns = e.get('functions', [])
            lo, hi = count_of(fns)
            chance = 1.0; player = False
            for c in pc + e.get('conditions', []):
                if c['condition'] in ('minecraft:random_chance', 'minecraft:random_chance_with_enchanted_bonus'):
                    chance = c.get('chance', c.get('unenchanted_chance', 1.0))
                if c['condition'] == 'minecraft:killed_by_player': player = True
            out.append(dict(mob=mob, item=e['name'], item_name=item_name(e['name']), min=lo, max=hi, chance=chance,
                            looting=any(f['function'] in ('minecraft:enchanted_count_increase', 'minecraft:looting_enchant') for f in fns),
                            cooks_if_on_fire=any(f['function'] == 'minecraft:furnace_smelt' for f in fns), player_kill_only=player))
    return out

# ------------------------------------------------------------------ spawns (pending biome modifiers)
spawns = []
for f in sorted(glob.glob(f'{REPO}/docs/zero-g-tweaks-bundle/pending-data/neoforge/biome_modifier/spawns_*.json')):
    o = json.load(open(f))
    for s in o['spawners']:
        spawns.append(dict(mob=s['type'].split(':')[1], biomes=' '.join(o['biomes']) if isinstance(o['biomes'], list) else o['biomes'],
                           weight=s['weight'], min_group=s['minCount'], max_group=s['maxCount'],
                           category=S[s['type'].split(':')[1]]['mob_category'], source=os.path.relpath(f, REPO)))

# ------------------------------------------------------------------ recipes: raw -> cooked -> dishes
REC = {}
for f in glob.glob(f'{RES}/data/zerog_tweaks/recipe/*.json'):
    REC[os.path.basename(f)[:-5]] = json.load(open(f))
def ingredients(r):
    if 'ingredient' in r: return [r['ingredient']]
    if 'ingredients' in r: return r['ingredients']
    if 'key' in r: return list(r['key'].values())
    return []
def items_in(r):
    out = set()
    for i in ingredients(r):
        for x in (i if isinstance(i, list) else [i]):
            if isinstance(x, dict) and 'item' in x: out.add(x['item'])
            elif isinstance(x, str): out.add(x)
    return out
def result(r):
    x = r.get('result'); return x['id'] if isinstance(x, dict) else x
COOKED, USED_IN = {}, {}
for n, r in REC.items():
    res = result(r)
    if not res: continue
    for i in items_in(r):
        if r['type'] in ('minecraft:smelting', 'minecraft:smoking', 'minecraft:campfire_cooking'):
            COOKED.setdefault(i, set()).add(res)
        elif r['type'].startswith('minecraft:crafting'):
            USED_IN.setdefault(i, set()).add(res)

# ------------------------------------------------------------------ build rows
def geo_bones(mob):
    j = json.load(open(f'{A}/geckolib/models/entity/{mob}.geo.json'))
    return [b['name'] for b in j['minecraft:geometry'][0].get('bones', [])]
def anim_names(mob):
    p = f'{A}/geckolib/animations/entity/{mob}.animation.json'
    return [a.split('.')[-1] for a in json.load(open(p))['animations']] if os.path.exists(p) else []

rows, drop_rows, anim_rows = [], [], []
for k in MOBS:
    d, s = DESIGN[k], S[k]
    st = dict(d['stats'])
    w, h, l = hitbox(d['info'])
    dr = drops(k); drop_rows += dr
    have = anim_names(k)
    need = [a.split(' ')[0] for a in s['animations']]
    for a in need:
        anim_rows.append(dict(mob=k, animation=f'animation.{k}.{a}', exists=a in have,
                              loop=a in ('idle', 'walk', 'run', 'swim', 'fly', 'float', 'glide', 'crawl', 'walk_sideways', 'dormant', 'burrow')))
    row = dict(
        id=f'zerog_tweaks:{k}', key=k, name=d['name'], sheet=SHEET[k], galaxy=s['galaxy'], world=s['world'], role=s['role'],
        dimensions=s['dimensions'], mob_category=s['mob_category'], java_base=s['java_base'],
        hitbox_width=w, hitbox_height=h, hitbox_length=l,
        max_health=first_int(st.get('Health')), attack_damage=first_int(st.get('Damage')) if st.get('Damage', 'None') != 'None' else 0,
        damage_text=st.get('Damage', 'None'), movement_speed=s['movement_speed'], speed_class=s['speed_class'], flying_speed=s['flying_speed'],
        follow_range=s['follow_range'], armor=s['armor'], knockback_resistance=s['knockback_resistance'],
        breeding_item=s['breeding_item'] and f"zerog_tweaks:{s['breeding_item']}", interactions=s['interactions'], immunities=s['immunities'],
        behavior=st.get('Behavior') or st.get('Phases'), spawns_text=st.get('Spawns') or st.get('Arena'),
        drops_text=st.get('Drops'), design_notes=d['notes'] + [f'{a}: {b}' for a, b in d['stats']
                                                             if a not in ('Health', 'Damage', 'Speed', 'Behavior', 'Spawns', 'Drops', 'Breeding', 'Phases', 'Arena')],
        goals_in_priority_order=s['goals'], states=s['states'], transitions=s['transitions'], animations_required=[f'animation.{k}.{a}' for a in need],
        animations_existing=[f'animation.{k}.{a}' for a in have], mechanics=s['mechanics'], boss=s['boss'],
        vanilla_base=s['vanilla']['base'], vanilla_also=s['vanilla'].get('also', []), vanilla_rig=s['vanilla']['rig'],
        animation_references={f'animation.{k}.{a}': ref for a, ref in s['vanilla']['anims'].items()},
        loot_table=f'zerog_tweaks:entities/{k}', drops=[{x: y for x, y in r.items() if x != 'mob'} for r in dr],
        spawn_egg={'base': eggs.get(k, ('', ''))[0], 'spots': eggs.get(k, ('', ''))[1], 'item': f'zerog_tweaks:{k}_spawn_egg'},
        files={'geo_model': f'assets/zerog_tweaks/geckolib/models/entity/{k}.geo.json',
               'animations': f'assets/zerog_tweaks/geckolib/animations/entity/{k}.animation.json',
               'texture': f'assets/zerog_tweaks/textures/entity/{k}.png',
               'glowmask': f'assets/zerog_tweaks/textures/entity/{k}_glowmask.png',
               'reference_sheet': f'docs/zero-g-tweaks-bundle/sheets/mobs/{SHEET[k]}'},
        geo_bones=geo_bones(k),
        spawns=[{x: y for x, y in r.items() if x != 'mob'} for r in spawns if r['mob'] == k])
    if not os.path.exists(f"{RES}/{row['files']['glowmask']}"): row['files']['glowmask'] = None
    for p in row['files'].values():
        assert p is None or os.path.exists(f'{RES}/{p}') or os.path.exists(f'{REPO}/{p}'), p
    rows.append(row)

json.dump({'schema': 'zerog_tweaks.mobs.v1', 'count': len(rows), 'mobs': rows}, open(f'{DATA}/mobs.json', 'w'), indent=2)

FLAT = ['key', 'name', 'galaxy', 'world', 'role', 'mob_category', 'java_base', 'vanilla_base', 'hitbox_width', 'hitbox_height', 'hitbox_length', 'max_health',
        'attack_damage', 'damage_text', 'movement_speed', 'flying_speed', 'follow_range', 'armor', 'knockback_resistance', 'breeding_item',
        'immunities', 'behavior', 'spawns_text', 'drops_text', 'loot_table', 'sheet']
with open(f'{DATA}/mobs.csv', 'w', newline='') as f:
    wr = csv.DictWriter(f, FLAT); wr.writeheader()
    for r in rows: wr.writerow({k: ('; '.join(r[k]) if isinstance(r[k], list) else '' if r[k] is None else r[k]) for k in FLAT})
def dump(name, rs, cols):
    with open(f'{DATA}/{name}', 'w', newline='') as f:
        wr = csv.DictWriter(f, cols); wr.writeheader(); [wr.writerow({c: r.get(c, '') for c in cols}) for r in rs]
dump('mob_drops.csv', drop_rows, ['mob', 'item', 'item_name', 'min', 'max', 'chance', 'looting', 'cooks_if_on_fire', 'player_kill_only'])
dump('mob_spawns.csv', spawns, ['mob', 'category', 'weight', 'min_group', 'max_group', 'biomes', 'source'])
dump('mob_animations.csv', anim_rows, ['mob', 'animation', 'exists', 'loop'])
dump('mob_state_transitions.csv', [dict(mob=r['key'], **t) for r in rows for t in r['transitions']], ['mob', 'frm', 'to', 'trigger'])
dump('mob_vanilla_bases.csv', [dict(mob=r['key'], name=r['name'], vanilla_base=r['vanilla_base'], also=' '.join(r['vanilla_also']), rig=r['vanilla_rig']) for r in rows],
     ['mob', 'name', 'vanilla_base', 'also', 'rig'])
dump('mob_animation_references.csv', [dict(mob=r['key'], animation=a, copy_from=ref, exists=a in r['animations_existing']) for r in rows for a, ref in r['animation_references'].items()],
     ['mob', 'animation', 'copy_from', 'exists'])
dump('mob_ai_goals.csv', [dict(mob=r['key'], priority=i, goal=g) for r in rows for i, g in enumerate(r['goals_in_priority_order'])], ['mob', 'priority', 'goal'])
attr_rows = [dict(mob=r['key'], **{'minecraft:generic.max_health': r['max_health'], 'minecraft:generic.attack_damage': r['attack_damage'],
              'minecraft:generic.movement_speed': r['movement_speed'], 'minecraft:generic.flying_speed': r['flying_speed'] or '',
              'minecraft:generic.follow_range': r['follow_range'], 'minecraft:generic.armor': r['armor'],
              'minecraft:generic.knockback_resistance': r['knockback_resistance']}) for r in rows]
dump('mob_attributes.csv', attr_rows, list(attr_rows[0]))

# ------------------------------------------------------------------ diagrams (Mermaid source + rendered PNG)
def safe(s): return re.sub(r'[^A-Za-z0-9_]', '_', s)
def label(s): return s.replace('"', "'")
WORLDS = [('Sol · Moon', 'Moon'), ('Sol · Mars', 'Mars'), ('Galaxy 2 · Cerulon', 'Cerulon'), ('Galaxy 3 · Skarn', 'Skarn'),
          ('Galaxy 4 · Eidolon', 'Eidolon'), ('Galaxy 5 · Solvane', 'Solvane')]
ROLE_CLASS = {'passive': 'passive', 'neutral': 'neutral', 'hostile': 'hostile', 'mini_boss': 'boss', 'boss': 'boss', 'final_boss': 'boss'}
CLASSDEF = ['classDef passive fill:#d9f2e3,stroke:#2e7d4f,color:#123', 'classDef neutral fill:#fff1cc,stroke:#b8860b,color:#321',
            'classDef hostile fill:#ffd9d6,stroke:#b3261e,color:#300', 'classDef boss fill:#e6dcff,stroke:#5b3fb3,color:#102,stroke-width:2px']

def world_map():
    L = ['flowchart LR']
    # Mermaid stacks LR subgraphs bottom-up, so list them in reverse to read Moon -> bosses top-down
    groups = list(reversed(WORLDS + [('Wastelands · every galaxy (tinted)', 'wasteland'), ('Guardian bosses · key worlds', 'key')]))
    for title, w in groups:
        L.append(f'  subgraph {safe(title)}["{title}"]')
        L.append('    direction TB')
        for r in rows:
            match = (w == 'wasteland' and 'wasteland' in r['world']) or (w == 'key' and r['role'] in ('boss', 'final_boss')) or \
                    (w not in ('wasteland', 'key') and r['world'] == w and r['role'] not in ('boss', 'final_boss'))
            if match:
                extra = f"<br/>{r['world'].replace(' wasteland', '')}" if w == 'wasteland' else ''
                L.append(f'    {r["key"]}["{r["name"]}<br/>{r["max_health"]} hp{extra}"]:::{ROLE_CLASS[r["role"]]}')
        L.append('  end')
    return '\n'.join(L + ['  ' + c for c in CLASSDEF])

FOODS = {'zerog_tweaks:' + n.lower() for n in re.findall(r'FoodProperties (\w+) =', open(f'{REPO}/src/main/java/net/zerog/tweaks/item/ZGFoods.java').read())}
def food_chain():
    L = ['flowchart LR']
    COOK = {k: {c for c in v if c in FOODS} for k, v in COOKED.items()}
    USE = {k: {c for c in v if c in FOODS} for k, v in USED_IN.items()}
    seen = set()
    for r in rows:
        for d in r['drops']:
            it = d['item']
            if not (COOK.get(it) or USE.get(it) or it in FOODS): continue
            L.append(f'  {r["key"]}(["{r["name"]}"]):::{ROLE_CLASS[r["role"]]} --> {safe(it)}["{label(d["item_name"])}"]')
            for c in sorted(COOK.get(it, [])):
                L.append(f'  {safe(it)} -- cook --> {safe(c)}["{label(item_name(c))}"]:::cooked')
                for dish in sorted(USE.get(c, [])):
                    if (c, dish) in seen: continue
                    seen.add((c, dish)); L.append(f'  {safe(c)} --> {safe(dish)}["{label(item_name(dish))}"]:::dish')
            for dish in sorted(USE.get(it, [])):
                if (it, dish) in seen: continue
                seen.add((it, dish)); L.append(f'  {safe(it)} --> {safe(dish)}["{label(item_name(dish))}"]:::dish')
    return '\n'.join(L + ['  ' + c for c in CLASSDEF] + ['  classDef cooked fill:#ffe6cc,stroke:#c26a00,color:#310',
                                                         '  classDef dish fill:#e0ecff,stroke:#2458b3,color:#012'])

def boss_chain():
    return '\n'.join(['flowchart LR',
        '  T2["T2 gate<br/>(Sol materials)"] --> G2(("Galaxy 2"))',
        '  G2 --> PS["Prism Sentinel<br/>300 hp · 3 phases"]:::boss',
        '  PS --> K3["Galaxy 3 Gate Key"] & CT["Cerulite template"] & SP["Sentinel Prism<br/>T3 lens upgrade"]',
        '  K3 --> G3(("Galaxy 3")) --> RT["Rift Tyrant<br/>450 hp · 3 phases"]:::boss',
        '  RT --> K4["Galaxy 4 Gate Key"] & ST["Skarnite template"] & RH["Rift Heart"]',
        '  K4 --> G4(("Galaxy 4")) --> EC["Eidolon Captain<br/>400 hp · fight or parley"]:::boss',
        '  EC --> K5["Galaxy 5 Gate Key"] & ET["Eidolite template"] & CL["Captain\'s Lantern<br/>(if spared)"]',
        '  G4 -.-> FW["Frost Warden<br/>mini-boss · wrecks"]:::boss --> RS["Remnant Shards<br/>(parley needs 16)"] -.-> EC',
        '  K5 --> G5(("Galaxy 5")) --> SC["Sun Colossus<br/>mini-boss · plateaus"]:::boss --> SV["Solvanite template<br/>Colossus Core"]',
        '  G5 --> DS["The Dying Star<br/>900 hp · final boss"]:::boss --> HS["Heart of Solvane<br/>Rekindle or Let it fade"]',
        ] + ['  ' + c for c in CLASSDEF])

def state_diagram(r):
    st = r['states']
    ids = [f's{i}' for i in range(len(st))]
    L = ['stateDiagram-v2', '  direction LR']
    for i, s in zip(ids, st): L.append(f'  state "{label(s)}" as {i}')
    L.append(f'  [*] --> {ids[0]}')
    for a, b, t in S[r['key']]['_T']:
        dst = '[*]' if b == -1 else ids[b]
        L.append(f'  {ids[a]} --> {dst}' + (f' : {label(t).replace(":", " -")}' if t else ''))
    return '\n'.join(L)

PP = f'{OUT}/.pp.json'
json.dump({'executablePath': '/opt/pw-browsers/chromium', 'args': ['--no-sandbox']}, open(PP, 'w'))
def render(name, src, width=2400):
    mmd = f'{DIA}/{name}.mmd'; open(mmd, 'w').write(src + '\n')
    subprocess.run(['mmdc', '-p', PP, '-i', mmd, '-o', f'{DIA}/{name}.png', '-w', str(width), '-b', 'white', '-q'], check=True,
                   stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)

DIAGRAMS = {'mob_world_map': world_map(), 'mob_drops_food_chain': food_chain(), 'boss_progression': boss_chain()}
for n, src in DIAGRAMS.items(): render(n, src)
for r in rows: render(f'states/{r["key"]}', state_diagram(r), 1400)
os.remove(PP)

# one Markdown page with every diagram inline (GitHub renders ```mermaid blocks)
md = ['# Mob diagrams', '', 'Mermaid sources render on GitHub. PNG copies sit next to each `.mmd` file.', '']
for n, src in DIAGRAMS.items():
    md += [f'## {n.replace("_", " ").capitalize()}', '', '```mermaid', src, '```', '']
md += ['## Behavior states per mob', '']
for r in rows:
    md += [f'### {r["name"]}', '', '```mermaid', state_diagram(r), '```', '']
open(f'{DIA}/README.md', 'w').write('\n'.join(md))
print('mobs', len(rows), 'drops', len(drop_rows), 'spawns', len(spawns), 'anims', len(anim_rows),
      'missing anims', sum(not a['exists'] for a in anim_rows))

# ------------------------------------------------------------------ workbook for humans (same data as the CSVs)
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
wb = Workbook(); wb.remove(wb.active)
for name in ['mobs', 'mob_vanilla_bases', 'mob_animation_references', 'mob_attributes', 'mob_drops', 'mob_spawns', 'mob_ai_goals', 'mob_state_transitions', 'mob_animations']:
    ws = wb.create_sheet(name)
    for i, row in enumerate(csv.reader(open(f'{DATA}/{name}.csv'))):
        ws.append([float(c) if re.fullmatch(r'-?\d+(\.\d+)?', c) else c for c in row])
    for c in ws[1]:
        c.font = Font(bold=True, color='FFFFFF'); c.fill = PatternFill('solid', fgColor='2F3E57'); c.alignment = Alignment(vertical='center')
    ws.freeze_panes = 'B2'; ws.auto_filter.ref = ws.dimensions
    for col in ws.columns:
        ws.column_dimensions[col[0].column_letter].width = min(60, max(10, max(len(str(c.value or '')) for c in col) + 2))
wb.save(f'{DATA}/mobs.xlsx')
